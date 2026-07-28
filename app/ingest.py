# -*- coding: utf-8 -*-
"""Excel 적재 및 검증.

공통 규칙
- 'DATA' 시트를 읽고 1행을 헤더로 사용.
- TM-NO는 문자열로만 취급. 셀이 날짜(datetime)로 변환되어 있으면 오류로 거부.
- 마스터는 전체 교체(idempotent). 트랜잭션은 삽입(UNIQUE는 UPSERT).
"""
import calendar
import csv
import datetime
import io
from openpyxl import load_workbook

from .calc import parse_alloc_rule, base_tmno, largest_remainder
from . import db

# CSV 제품목록의 공정 컬럼 (좌→우 = 공정 순서)
CSV_PROC_COLS = ["성형", "소결", "정형", "가공", "압입", "밴딩", "후처리"]


class IngestError(Exception):
    pass


def _rows(path, tm_col=None):
    """DATA 시트를 [dict,...] 로. tm_col(헤더명) 지정 시 문자열 강제/날짜검증."""
    wb = load_workbook(path, data_only=True)
    ws = wb["DATA"] if "DATA" in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else "" for h in next(it)]
    out = []
    for i, raw in enumerate(it, start=2):
        if all(c is None or str(c).strip() == "" for c in raw):
            continue
        row = {headers[j]: raw[j] for j in range(len(headers))}
        if tm_col is not None:
            v = row.get(tm_col)
            if isinstance(v, (datetime.date, datetime.datetime)):
                raise IngestError(
                    f"{i}행 TM-NO가 날짜로 변환됨({v}). 원본 셀을 텍스트로 고쳐 다시 올려주세요.")
            row[tm_col] = "" if v is None else str(v).strip()
        out.append((i, row))
    return headers, out


def _d(v):
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def _ym(v):
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v.strftime("%Y-%m")
    return str(v).strip()[:7]


def _int(v):
    return int(round(float(v))) if v not in (None, "") else 0


# ── 마스터 ─────────────────────────────────────────────
def ingest_process_master(conn, path):
    _, rows = _rows(path)
    data = []
    for i, r in rows:
        part = str(r.get("파트구분", "")).strip()
        name = str(r.get("공정명", "")).strip()
        excl = 1 if str(r.get("COPQ제외여부", "")).strip().upper() == "Y" else 0
        if not part or not name:
            raise IngestError(f"{i}행: 파트구분/공정명 필수")
        data.append((part, name, excl))
    conn.execute("DELETE FROM process")
    conn.executemany("INSERT INTO process(part,name,copq_exclude) VALUES(?,?,?)", data)
    conn.commit()
    return len(data)


def ingest_product_master(conn, path):
    _, rows = _rows(path, tm_col="TM-NO")
    prods, routes = {}, []
    for i, r in rows:
        tm = r["TM-NO"]
        name = str(r.get("품명", "")).strip()
        part = str(r.get("파트구분", "")).strip()
        seq = _int(r.get("공정순서"))
        proc = str(r.get("공정명", "")).strip()
        price = float(r.get("기준단가(원)") or 0)
        if not tm or not proc:
            raise IngestError(f"{i}행: TM-NO/공정명 필수")
        if not part:
            raise IngestError(f"{i}행: 파트구분 필수 (VMS PART/TM PART)")
        prods[tm] = (name, part)               # 파트는 제품 마스터가 직접 명시
        routes.append((tm, seq, proc, price))
    conn.execute("DELETE FROM product_route")
    conn.execute("DELETE FROM product")
    for tm, (name, part) in prods.items():
        conn.execute("INSERT INTO product(tm_no,name,part) VALUES(?,?,?)", (tm, name, part))
    conn.executemany("INSERT INTO product_route(tm_no,seq,process,unit_price) VALUES(?,?,?,?)", routes)
    conn.commit()
    return len(prods)


def ingest_product_csv(conn, path, part):
    """제품목록 CSV(wide, cp949) 적재. part = 'VMS PART' | 'TM PART'.
    컬럼: TM-NO, 품명, 중량, 성형, 소결, 정형, 가공, 압입, 밴딩, 후처리
    공정 칸에 값이 있으면 라우팅에 포함(좌→우 순서), 값은 op_code로 보존.
    반환: (제품수, 신규수)."""
    raw = open(path, "rb").read()
    text = None
    for enc in ("cp949", "euc-kr", "utf-8-sig", "utf-8"):
        try:
            text = raw.decode(enc); break
        except Exception:
            continue
    if text is None:
        raise IngestError("CSV 인코딩을 인식할 수 없습니다 (cp949/utf-8).")
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        raise IngestError("빈 CSV 파일")
    header = [h.strip() for h in rows[0]]
    idx = {h: i for i, h in enumerate(header)}
    if "TM-NO" not in idx or "품명" not in idx:
        raise IngestError("헤더에 TM-NO / 품명 이 필요합니다.")
    db.ensure_processes(conn, part)

    existing = {r["tm_no"] for r in conn.execute("SELECT tm_no FROM product WHERE part=?", (part,))}
    n = new = 0
    for r in rows[1:]:
        if not r or not str(r[idx["TM-NO"]]).strip():
            continue
        tm = base_tmno(r[idx["TM-NO"]])       # 변형(598-10A)→base(598-10) 합침
        name = str(r[idx["품명"]]).strip() if idx.get("품명", -1) < len(r) else ""
        weight = 0.0
        if "중량" in idx and idx["중량"] < len(r):
            try:
                weight = float(str(r[idx["중량"]]).strip() or 0)
            except ValueError:
                weight = 0.0
        route = []
        for proc in CSV_PROC_COLS:
            if proc in idx and idx[proc] < len(r):
                code = str(r[idx[proc]]).strip()
                if code:
                    route.append((proc, code))
        conn.execute("INSERT INTO product(tm_no,name,part,weight) VALUES(?,?,?,?) "
                     "ON CONFLICT(tm_no) DO UPDATE SET name=excluded.name, part=excluded.part, "
                     "weight=excluded.weight", (tm, name, part, weight))
        conn.execute("DELETE FROM product_route WHERE tm_no=?", (tm,))
        for seq, (proc, code) in enumerate(route, start=1):
            conn.execute("INSERT INTO product_route(tm_no,seq,process,unit_price,op_code) "
                         "VALUES(?,?,?,?,?)", (tm, seq, proc, 0, code))
        n += 1
        if tm not in existing:
            new += 1
    conn.commit()
    return n, new


def _clean_name(v):
    """헤더/불량명 정규화: 개행 제거·공백 정리."""
    return " ".join(str(v or "").replace("\n", " ").split())


# 불량유형 마스터 시트 → 파트 매핑
DEFECT_MASTER_SHEETS = {"DATA_1PART": "VMS PART", "DATA_2PART": "TM PART"}


def ingest_defect_master(conn, path):
    """불량유형 마스터(DATA_1PART/DATA_2PART 2시트) 적재.
    - 구분(공정/셋팅), 발생공정(공란 허용=입력공정 귀속), 불량명, 배분기준.
    - 발생공정 지정 시 집계공정(성형·소결·정형·가공·기타)만 허용.
    반환: (건수, errors[list]). errors 있으면 반영하지 않음."""
    wb = load_workbook(path, data_only=True)
    agg = set(db.AGG_PROCESSES)
    data, errors, seen = [], [], set()
    sheets = [s for s in DEFECT_MASTER_SHEETS if s in wb.sheetnames]
    if not sheets:
        return 0, ["DATA_1PART/DATA_2PART 시트를 찾을 수 없습니다."]
    for sheet in sheets:
        part = DEFECT_MASTER_SHEETS[sheet]
        ws = wb[sheet]
        it = ws.iter_rows(values_only=True)
        headers = [_clean_name(h) for h in next(it)]
        idx = {h: j for j, h in enumerate(headers)}
        need = ["구분", "공정명(발생공정)", "불량명", "배분기준"]
        miss = [h for h in need if h not in idx]
        if miss:
            errors.append(f"[{sheet}] 헤더 누락: {', '.join(miss)}")
            continue
        for i, raw in enumerate(it, start=2):
            if all(c is None or str(c).strip() == "" for c in raw):
                continue
            kind = str(raw[idx["구분"]] or "").strip()
            proc = _clean_name(raw[idx["공정명(발생공정)"]])
            name = _clean_name(raw[idx["불량명"]])
            rule = str(raw[idx["배분기준"]] or "").strip()
            if kind not in ("공정", "셋팅"):
                errors.append(f"[{sheet}] {i}행: 구분은 공정/셋팅만 ('{kind}')")
            if not name:
                errors.append(f"[{sheet}] {i}행: 불량명 필수")
                continue
            key = (part, kind, name)
            if key in seen:
                errors.append(f"[{sheet}] {i}행: 불량명 중복 ('{name}')")
            seen.add(key)
            if proc and proc not in agg:
                errors.append(f"[{sheet}] {i}행: 발생공정 '{proc}'은 집계공정(성형·소결·정형·가공·기타)만 허용")
            for p, _rt in parse_alloc_rule(rule):
                if p not in agg:
                    errors.append(f"[{sheet}] {i}행: 배분기준 공정 '{p}'은 집계공정만 허용")
            data.append((part, kind, proc, name, rule))
    if errors:
        return 0, errors
    conn.execute("DELETE FROM defect_type")
    conn.executemany(
        "INSERT INTO defect_type(part,kind,process,name,alloc_rule) VALUES(?,?,?,?,?)", data)
    conn.commit()
    return len(data), []


# ── 100EA 검토(외주/폐기 공통) ──────────────────────────
def _reingest_preserve_reviewed(conn, source, rows):
    """재스캔 idempotent 적재이면서, 사람이 검토(승인/수정/반려)한 행과 **이번 달 이전(지난달까지) 행**은
    원본이 바뀌어도 건드리지 않고 그대로 보존한다(마감된 과거월 데이터 잠금).

    rows: [(d,tm_no,defect_name,qty,part,proc,kind,status,reg_user), ...] (source는 고정값으로 별도 처리)
    이번 달 1일 이전 날짜의 신규행은 통째로 무시(DB에 이미 있는 값 유지), reviewed=1인 기존 행과
    (d,tm_no,defect_name) 키가 같은 신규행도 버린다. reviewed=0이면서 이번 달 이후인 기존 행만
    삭제 후 나머지를 재적재한다.
    반환: (적재건수, 검토완료라 건너뛴 건수, 과거월 잠금이라 건너뛴 건수)."""
    cur_month_start = datetime.date.today().replace(day=1).isoformat()
    reviewed_keys = {(r["d"], r["tm_no"], r["defect_name"]) for r in conn.execute(
        "SELECT d,tm_no,defect_name FROM defect_entry WHERE source=? AND reviewed=1", (source,))}
    conn.execute("DELETE FROM defect_entry WHERE source=? AND reviewed=0 AND d>=?", (source, cur_month_start))
    keep, skipped, locked = [], 0, 0
    for row in rows:
        d, tm_no, defect_name = row[0], row[1], row[2]
        if d < cur_month_start:
            locked += 1
            continue
        if (d, tm_no, defect_name) in reviewed_keys:
            skipped += 1
            continue
        keep.append(row)
    conn.executemany(
        "INSERT INTO defect_entry(d,tm_no,defect_name,qty,part,process,kind,source,status,reg_user) "
        "VALUES(?,?,?,?,?,?,?,?,?,?)", [(*row[:7], source, row[7], row[8]) for row in keep])
    conn.commit()
    return len(keep), skipped, locked


# ── 사내불량 일일 양식(폴더 스캔) ──────────────────────
import os
import re

# 파일명: {PART}_{공정}_{구분}불량_{YYYYMMDD}.xlsx
_DAILY_RE = re.compile(r"^(1PART|2PART)_(.+?)_(공정|셋팅)불량_(\d{8})$")
_PART_TOKEN = {"1PART": "VMS PART", "2PART": "TM PART"}
_DAILY_PROCS = {
    "VMS PART": {"성형", "소결", "정형", "가공", "압입", "밴딩", "선별공정"},
    "TM PART": {"성형", "소결", "정형", "가공", "선별공정"},
}


def parse_daily_filename(fname):
    """'1PART_성형_공정불량_20260724.xlsx' → (part, proc, kind, date, batch_key) 또는 None."""
    stem = os.path.splitext(os.path.basename(fname))[0]
    mo = _DAILY_RE.match(stem)
    if not mo:
        return None
    part = _PART_TOKEN[mo.group(1)]
    proc, kind, ymd = mo.group(2), mo.group(3), mo.group(4)
    d = f"{ymd[:4]}-{ymd[4:6]}-{ymd[6:8]}"
    return part, proc, kind, d, f"{part}|{proc}|{kind}|{d}"


def _find_header_row(ws, key="TM-NO"):
    for r in range(1, min(ws.max_row, 15) + 1):
        for c in range(1, min(ws.max_column, 10) + 1):
            if str(ws.cell(row=r, column=c).value or "").strip() == key:
                return r
    return None


def ingest_daily_defect_file(conn, path, user=""):
    """사내불량 일일 양식 1개 적재. 행=TM-NO, 열=불량유형, 셀=수량.

    공정(성형시트가 어디서 작성됐는지)은 그대로 원본 그대로 저장만 하고,
    '이 불량이 어느 집계공정 몫인지'는 여기서 정하지 않는다 — calc.Masters.resolve()가
    마스터의 발생공정/배분기준을 보고 최종 결정한다(공란/입력공정100%인 유형만 여기 저장한
    시트 공정을 그대로 씀). 성형 시트는 애초에 불량 D/B에 반영하지 않는다(전량 제외).
    같은 batch_key(파트|공정|구분|일자)는 삭제 후 재적재(idempotent).
    반환: (적재건수, errors[list])."""
    parsed = parse_daily_filename(path)
    if not parsed:
        return 0, [f"파일명 형식 오류: {os.path.basename(path)} "
                   f"(예: 1PART_성형_공정불량_20260724.xlsx)"]
    part, sheet_proc, kind, d, batch_key = parsed
    if sheet_proc not in _DAILY_PROCS.get(part, set()):
        return 0, [f"{os.path.basename(path)}: '{sheet_proc}'은 {part}의 공정이 아님"]
    if sheet_proc == "성형":
        # 성형 공정에서 작성한 불량시트(공정/셋팅 모두)는 D/B에 반영하지 않는다.
        conn.execute("DELETE FROM defect_entry WHERE batch_key=?", (batch_key,))
        conn.commit()
        return 0, []

    # 불량유형이 마스터(part,kind)에 등록돼 있는지만 확인(발생공정 판단은 calc가 함)
    names = {r["name"] for r in conn.execute(
        "SELECT name FROM defect_type WHERE part=? AND kind=?", (part, kind))}

    wb = load_workbook(path, data_only=True)
    ws = wb["DATA"] if "DATA" in wb.sheetnames else wb.worksheets[0]
    hr = _find_header_row(ws)
    if hr is None:
        return 0, [f"{os.path.basename(path)}: 헤더행(TM-NO)을 찾을 수 없음"]
    headers = [str(ws.cell(row=hr, column=c).value or "").strip()
               for c in range(1, ws.max_column + 1)]
    tm_col = headers.index("TM-NO") + 1
    # 불량유형 열 = TM-NO/품명/No 이후의 헤더
    fixed = {"No", "TM-NO", "품명", ""}
    dcols = [(c + 1, headers[c]) for c in range(len(headers)) if headers[c] not in fixed]

    rows, errors = [], []
    for r in range(hr + 1, ws.max_row + 1):
        tmv = ws.cell(row=r, column=tm_col).value
        if isinstance(tmv, (datetime.date, datetime.datetime)):
            errors.append(f"{os.path.basename(path)} {r}행: TM-NO가 날짜로 변환됨({tmv})")
            continue
        tmno = base_tmno(tmv)
        if not tmno or tmno == "합계":
            continue
        for c, name in dcols:
            qty = _int(ws.cell(row=r, column=c).value)
            if qty <= 0:
                continue
            if name not in names:
                errors.append(f"{os.path.basename(path)}: 불량유형 '{name}'이 {part} {kind} 마스터에 없음")
                continue
            # 공정은 이 시트(sheet_proc)를 그대로 저장 — 최종 귀속은 calc.resolve()가 결정
            rows.append((d, tmno, name, qty, part, sheet_proc, kind, "direct", user, batch_key))
    if errors:
        return 0, errors
    conn.execute("DELETE FROM defect_entry WHERE batch_key=?", (batch_key,))
    conn.executemany(
        "INSERT INTO defect_entry(d,tm_no,defect_name,qty,part,process,kind,source,status,reg_user,batch_key) "
        "VALUES(?,?,?,?,?,?,?,?, 'confirmed', ?,?)", rows)
    conn.commit()
    return len(rows), []


def ingest_daily_defect_folder(conn, folder, user=""):
    """폴더(하위 포함) 사내불량 일일 양식 전체 스캔 적재.

    폴더 구조(년월/파트/공정)는 사람이 찾기 위한 것이고, 시스템은 **파일명**으로
    (파트·공정·구분·일자)를 판단한다 → 하위 폴더를 재귀 탐색하고 위치는 따지지 않는다.
    Excel 임시파일(~$…)은 건너뛴다.
    반환: (파일수, 총건수, 상세[(상대경로, 건수, errors)])."""
    detail, files, total = [], 0, 0
    for root, _dirs, names in os.walk(folder):
        for fn in sorted(names):
            if fn.startswith("~$") or not fn.lower().endswith((".xlsx", ".xlsm")):
                continue
            if not _DAILY_RE.match(os.path.splitext(fn)[0]):
                continue
            path = os.path.join(root, fn)
            n, errs = ingest_daily_defect_file(conn, path, user)
            files += 1
            total += n
            detail.append((os.path.relpath(path, folder), n, errs))
    return files, total, detail


# ── 사내불량 월별 양식(시트=일자 1~31 + 종합) ──────────────
# 파일명: {YYMM}_{공정}_{공정불량|셋팅불량}_생산{1|2}파트.xlsx  (예: 2607_성형_공정불량_생산1파트.xlsx)
_MONTHLY_DEFECT_RE = re.compile(r"^(\d{4})_(.+?)_(공정|셋팅)불량_생산([12])파트$")
_MONTHLY_PART_TOKEN = {"1": "VMS PART", "2": "TM PART"}
# 불량유형 열 중 집계 제외(폐기·소재불량은 폐기불량/외주소재불량 소스와 중복이라 별도 소스로만 집계, 소계는 합계열)
_MONTHLY_SKIP_COLS = {"폐기", "소재불량", "소계"}


def parse_monthly_defect_filename(fname):
    """'2607_성형_공정불량_생산1파트.xlsx' → (year,month,part,proc,kind) 또는 None."""
    stem = os.path.splitext(os.path.basename(fname))[0]
    mo = _MONTHLY_DEFECT_RE.match(stem)
    if not mo:
        return None
    yymm, proc, kind, partnum = mo.groups()
    year, month = 2000 + int(yymm[:2]), int(yymm[2:])
    if not (1 <= month <= 12):
        return None
    return year, month, _MONTHLY_PART_TOKEN[partnum], proc, kind


def ingest_monthly_defect_file(conn, path, user=""):
    """사내불량 월별 양식 1개 적재. 시트 `1`~해당월 마지막일=일자별 데이터, `종합`은 무시.
    행=TM-NO, 열=세부 불량유형(4행 헤더), 셀=수량.

    성형 공정에서 작성한 파일 규칙:
    - 구분=공정불량 → D/B에 전혀 반영하지 않는다(파일 통째로 skip).
    - 구분=셋팅불량 → D/B엔 반영해 셋팅불량율에는 포함시키되, Scrap Cost·COPQ 계산에서는
      제외한다(defect_entry.exclude_cost=1).
    같은 batch_key(파트|공정|구분|일자)는 일자별로 삭제 후 재적재(idempotent, 기존 일일양식과 동일 방식).
    반환: (적재건수, errors[list])."""
    parsed = parse_monthly_defect_filename(path)
    if not parsed:
        return 0, [f"파일명 형식 오류: {os.path.basename(path)} "
                   f"(예: 2607_성형_공정불량_생산1파트.xlsx)"]
    year, month, part, proc, kind = parsed
    if proc == "성형" and kind == "공정":
        return 0, []                                # 성형 작성 공정불량은 D/B 전량 미반영
    exclude_cost = 1 if (proc == "성형" and kind == "셋팅") else 0

    names = {r["name"] for r in conn.execute(
        "SELECT name FROM defect_type WHERE part=? AND kind=?", (part, kind))}

    wb = load_workbook(path, data_only=True)
    days_in_month = calendar.monthrange(year, month)[1]
    rows, errors = [], []

    for day in range(1, days_in_month + 1):
        sheet_name = str(day)
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        d = f"{year:04d}-{month:02d}-{day:02d}"
        batch_key = f"{part}|{proc}|{kind}|{d}"

        tm_col = None
        for c in range(1, ws.max_column + 1):
            if str(ws.cell(row=5, column=c).value or "").strip() == "TM NO.":
                tm_col = c
                break
        if tm_col is None:
            errors.append(f"{os.path.basename(path)} [{sheet_name}]: 헤더(TM NO.)를 찾을 수 없음")
            continue
        # 세부 불량유형 헤더는 4행. 공백/줄바꿈을 모두 제거해 마스터 불량명과 비교한다
        # (예: xlsm의 "영점 및\n동심도수정" → "영점및동심도수정").
        dcols = []
        for c in range(1, ws.max_column + 1):
            raw = str(ws.cell(row=4, column=c).value or "")
            name = "".join(raw.split())
            if not name or name in _MONTHLY_SKIP_COLS or name not in names:
                continue
            dcols.append((c, name))

        for r in range(6, ws.max_row + 1):
            tmv = ws.cell(row=r, column=tm_col).value
            if isinstance(tmv, (datetime.date, datetime.datetime)):
                continue
            tmno = base_tmno(tmv)
            if not tmno:
                continue
            for c, name in dcols:
                qty = _int(ws.cell(row=r, column=c).value)
                if qty <= 0:
                    continue
                rows.append((d, tmno, name, qty, part, proc, kind, "direct",
                             "confirmed", user, batch_key, exclude_cost))
        conn.execute("DELETE FROM defect_entry WHERE batch_key=?", (batch_key,))

    conn.executemany(
        "INSERT INTO defect_entry(d,tm_no,defect_name,qty,part,process,kind,source,status,"
        "reg_user,batch_key,exclude_cost) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)", rows)
    conn.commit()
    return len(rows), errors


def ingest_monthly_defect_folder(conn, folder, user=""):
    """폴더(하위 포함) 사내불량 월별 양식 전체 스캔 적재. `01_사내불량\\{연도}\\{1파트|2파트}\\{월}\\`
    아래에 파일들이 섞여 있는 구조(하위폴더로 공정 구분 없음)를 그대로 재귀 탐색한다.
    Excel 임시파일(~$…)은 건너뛴다.
    반환: (파일수, 총건수, 상세[(상대경로, 건수, errors)])."""
    detail, files, total = [], 0, 0
    for root, _dirs, names in os.walk(folder):
        for fn in sorted(names):
            if fn.startswith("~$") or not fn.lower().endswith((".xlsx", ".xlsm")):
                continue
            if not _MONTHLY_DEFECT_RE.match(os.path.splitext(fn)[0]):
                continue
            path = os.path.join(root, fn)
            n, errs = ingest_monthly_defect_file(conn, path, user)
            files += 1
            total += n
            detail.append((os.path.relpath(path, folder), n, errs))
    return files, total, detail


# ── 트랜잭션 ───────────────────────────────────────────
def ingest_defect_entries(conn, path, source="direct", user="", quarantine_100=False):
    """불량입력/외주소재불량/폐기불량. quarantine_100: 단일TM×단일불량 ≥100 → pending."""
    _, rows = _rows(path, tm_col="TM-NO")
    n, pend = 0, 0
    for i, r in rows:
        tm = r["TM-NO"]; name = str(r.get("불량명", "")).strip(); qty = _int(r.get("수량"))
        if not tm or not name or qty <= 0:
            continue
        status = "pending" if (quarantine_100 and qty >= 100) else "confirmed"
        if status == "pending":
            pend += 1
        conn.execute("INSERT INTO defect_entry(d,tm_no,defect_name,qty,source,status,reg_user)"
                     " VALUES(?,?,?,?,?,?,?)",
                     (_d(r.get("일자")), tm, name, qty, source, status, user))
        n += 1
    conn.commit()
    return n, pend


def ingest_production(conn, path):
    _, rows = _rows(path, tm_col="TM-NO")
    n = 0
    for i, r in rows:
        tm = r["TM-NO"]; qty = _int(r.get("생산수량(EA)"))
        amt = float(r.get("생산금액(천원)") or 0)
        if not tm:
            continue
        conn.execute("INSERT INTO production(d,tm_no,qty,amount) VALUES(?,?,?,?) "
                     "ON CONFLICT(d,tm_no) DO UPDATE SET qty=excluded.qty, amount=excluded.amount",
                     (_d(r.get("일자")), tm, qty, amt))
        n += 1
    conn.commit()
    return n


import sqlite3

# 폐기 defect_process 중 집계버킷이 명확한 물리공정
_SCRAP_KNOWN_PROC = set(db.PHYS_TO_BUCKET)


_SCRAP_PART = {"1part": "VMS PART", "2part": "TM PART"}


def ingest_scrap_db(conn, path, quarantine_100=True):
    """폐기 scrap_data.db 적재. defect_category='해당' 만, defect_process→집계공정(수량 합산).
    폐기는 공정불량(kind='공정'). TM-NO 없는 공정단위 폐기(소결로_산화 등)는 part만으로 집계.
    불량명이 '기타_불량'이고 remark(비고)에 사유가 적혀있으면, '기타_불량' 대신 그 remark 내용을
    불량명으로 사용한다(TOP5 등 집계에도 세부 사유 그대로 반영).
    단일 TM×단일 불량명 ≥100 → pending(검토 대기, /input/outsource-review 화면에서 외주와 함께 관리).
    전체 재스캔 idempotent(source='discard' 삭제 후 재적재) — 사람이 검토(승인/수정/반려)한 행과
    이번 달 이전(지난달까지) 행은 재스캔해도 보존(잠금)한다.
    반환: (건수, note[dict])."""
    src = sqlite3.connect(path)
    src.row_factory = sqlite3.Row
    rows, unmapped = [], {}
    skip_noqty = skip_nopart = 0
    for r in src.execute("SELECT date,part,tmno,scrap_reason,quantity,defect_process,remark "
                         "FROM scrap_data WHERE defect_category='해당'"):
        part = _SCRAP_PART.get(str(r["part"] or "").strip().lower(), "")
        if not part:
            skip_nopart += 1
            continue
        qty = _int(r["quantity"])
        if qty <= 0:
            skip_noqty += 1                                 # 수량 없이 무게만 있는 폐기 → 스킵
            continue
        d = str(r["date"] or "")[:10]
        if not d:
            continue
        tmv = str(r["tmno"] or "").strip()
        tm = "" if tmv in ("", "-") else base_tmno(tmv)     # 제품 미지정 폐기는 tm 공란(파트로만 집계)
        raw_proc = str(r["defect_process"] or "").strip()
        proc = db.bucket_of(raw_proc)
        if raw_proc not in _SCRAP_KNOWN_PROC:
            unmapped[raw_proc] = unmapped.get(raw_proc, 0) + 1
            proc = "기타"                                   # 비표준 공정은 기타로(경고)
        reason = str(r["scrap_reason"] or "").strip() or "기타"
        remark = str(r["remark"] or "").strip()
        name = remark if (reason == "기타_불량" and remark) else reason
        status = "pending" if (quarantine_100 and qty >= 100) else "confirmed"
        rows.append((d, tm, name, qty, part, proc, "공정", status, ""))
    src.close()
    pend = sum(1 for row in rows if row[7] == "pending")
    n, _skipped, _locked = _reingest_preserve_reviewed(conn, "discard", rows)
    return n, {"수량없음_스킵": skip_noqty, "파트없음_스킵": skip_nopart, "비표준공정→기타": unmapped,
              "검토대기": pend}


_MISSING_PRICE_PART = {"1part": "VMS PART", "2part": "TM PART"}


def ingest_missing_price_xlsx(conn, path):
    """단가마스터에 단가가 아예 없던 TM-NO를 채우는 보정 파일 적재.
    헤더: TM-NO, 품명, PART, 성형단가, 소결단가, 정형단가, 기타단가(1행).
    정형단가가 0(또는 공란)이면 그 제품은 정형공정이 없는 것으로 보고 정형 행을 넣지 않는다.
    기타단가는 가공·기타 버킷 둘 다에 동일하게 저장한다(price_calc.py 규칙과 동일).
    제품 마스터에 없는 TM-NO는 건너뛴다. 반환: (반영 TM-NO수, note[dict])."""
    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    headers = [str(ws.cell(row=1, column=c).value or "").strip() for c in range(1, ws.max_column + 1)]
    need = ["TM-NO", "성형단가", "소결단가", "정형단가", "기타단가"]
    miss = [h for h in need if h not in headers]
    if miss:
        return 0, {"error": f"헤더 누락: {', '.join(miss)}"}
    idx = {h: headers.index(h) + 1 for h in need}

    rows, skipped_no_product, bad = [], [], 0
    for r in range(2, ws.max_row + 1):
        tmv = ws.cell(row=r, column=idx["TM-NO"]).value
        if tmv in (None, ""):
            continue
        tm = base_tmno(tmv)
        if not tm:
            continue
        if not conn.execute("SELECT 1 FROM product WHERE tm_no=?", (tm,)).fetchone():
            skipped_no_product.append(tm)
            continue
        try:
            sh = float(ws.cell(row=r, column=idx["성형단가"]).value or 0)
            so = float(ws.cell(row=r, column=idx["소결단가"]).value or 0)
            jh = float(ws.cell(row=r, column=idx["정형단가"]).value or 0)
            gi = float(ws.cell(row=r, column=idx["기타단가"]).value or 0)
        except (TypeError, ValueError):
            bad += 1
            continue
        rows.append((tm, "성형", sh))
        rows.append((tm, "소결", so))
        if jh:
            rows.append((tm, "정형", jh))
        rows.append((tm, "가공", gi))
        rows.append((tm, "기타", gi))
    conn.executemany(
        "INSERT INTO product_price(tm_no,process,unit_price,effective_from) VALUES(?,?,?,'2000-01-01') "
        "ON CONFLICT(tm_no,process,effective_from) DO UPDATE SET unit_price=excluded.unit_price",
        [(tm, proc, v) for tm, proc, v in rows])
    conn.commit()
    tm_count = len({tm for tm, _p, _v in rows})
    return tm_count, {"제품마스터없어_스킵": skipped_no_product, "비수치_무시": bad}


# 제품별 단가 Master(단가산출 시트) 블록: (블록명, TM열, 단가열)
_PRICE_BLOCKS = [("성형", 2, 4), ("소결", 6, 8), ("정형", 10, 12), ("완제품", 14, 16)]
_PRICE_SHEET = "단가산출"
_PRICE_HEADER_ROWS = 3          # 1~3행 헤더, 4행부터 데이터


def ingest_price_master(conn, path):
    """제품별 단가 Master 적재 → product_price(집계공정 5종 누적단가).

    규칙(사용자 확정):
      · 성형/소결/정형 = 각 블록 단가(누적)
      · 가공·기타 = **완제품 단가**
      · 정형이 'O'(정형공정 없음)이면 정형도 **완제품 단가** 적용
    반환: (품목수, note[dict])."""
    wb = load_workbook(path, data_only=True)
    if _PRICE_SHEET not in wb.sheetnames:
        return 0, {"error": f"'{_PRICE_SHEET}' 시트를 찾을 수 없음"}
    ws = wb[_PRICE_SHEET]

    block = {name: {} for name, _t, _p in _PRICE_BLOCKS}
    no_jeonghyeong = set()          # 정형 'O' = 정형공정 없음
    bad = 0
    for name, c_tm, c_pr in _PRICE_BLOCKS:
        for r in range(_PRICE_HEADER_ROWS + 1, ws.max_row + 1):
            tmv = ws.cell(row=r, column=c_tm).value
            if tmv in (None, ""):
                continue
            tm = base_tmno(tmv)
            if not tm:
                continue
            pv = ws.cell(row=r, column=c_pr).value
            if isinstance(pv, (int, float)):
                block[name][tm] = float(pv)
            elif name == "정형" and str(pv or "").strip().upper() == "O":
                no_jeonghyeong.add(tm)
            else:
                bad += 1            # #N/A 등

    tms = set(block["완제품"]) | set(block["성형"]) | set(block["소결"])
    rows, skipped = [], 0
    for tm in tms:
        final = block["완제품"].get(tm)
        if final is None:
            skipped += 1            # 완제품 단가 없으면 가공·기타를 정할 수 없음
            continue
        j = block["정형"].get(tm)
        if tm in no_jeonghyeong or j is None:
            j = final               # 정형공정 없음 → 이후는 완제품 단가
        for proc, price in (("성형", block["성형"].get(tm)), ("소결", block["소결"].get(tm)),
                            ("정형", j), ("가공", final), ("기타", final)):
            if price is None:
                continue
            rows.append((tm, proc, price))
    conn.executemany(
        "INSERT INTO product_price(tm_no,process,unit_price,effective_from) VALUES(?,?,?,'2000-01-01') "
        "ON CONFLICT(tm_no,process,effective_from) DO UPDATE SET unit_price=excluded.unit_price", rows)
    conn.commit()
    return len(tms) - skipped, {"정형없음(완제품단가)": len(no_jeonghyeong),
                                "완제품단가없어_스킵": skipped, "비수치_무시": bad}


_OUTSOURCE_PART = {"1part": "VMS PART", "2part": "TM PART"}

# 세부 불량명(xlsm 컬럼) → 집계공정. VBA의 "성형/소결 합계" 그룹 중 공정이 명확한 항목만 직접 매핑.
_OUTSOURCE_DETAIL_DIRECT = {
    "깨짐": "성형", "크랙": "성형", "기포": "성형", "뜯김": "성형", "밸런스NG": "성형",
    "변형(촥좌)": "소결", "변형(Window)": "소결", "변형(쏠림)": "소결", "브레이징": "소결",
    "조립불량": "소결", "변형(편가공)": "소결", "산화": "소결", "용침불량": "소결", "소결붙음": "소결",
}
# VBA에서 소결·기타에 50%씩 임의 분배하던 애매 항목 → TM-NO 라우팅에 정형이 있으면 정형, 없으면 소결.
_OUTSOURCE_DETAIL_AMBIGUOUS = ["찍힘.눌림", "이종혼입", "녹", "이물소착", "소재돌출"]
_OUTSOURCE_DETAIL_COLS = list(_OUTSOURCE_DETAIL_DIRECT) + _OUTSOURCE_DETAIL_AMBIGUOUS


def ingest_outsource_xlsm(conn, path, quarantine_100=True):
    """외주소재불량 xlsm(Sheet1) 적재. 파일이 미리 뭉쳐놓은 성형/소결/기타 합계 대신,
    19개 세부 불량명 컬럼을 그대로 불량명으로 반영해 TOP5 등에도 세부 사유가 나타나게 한다.
    애매 5종(찍힘.눌림/이종혼입/녹/이물소착/소재돌출)은 TM-NO의 제품 라우팅에 정형이 있으면 정형,
    없으면 소결로 배분. TM-NO 없는 행(제품 미지정, 날짜로 오인식된 경우 포함)은 라우팅 판단이
    불가능하므로 집계에서 제외한다. 공정불량(kind='공정').
    단일 TM×단일 불량명 ≥100 → pending(검토 대기, /input/outsource-review 화면).
    전체 재스캔 idempotent(source='outsource' 삭제 후 재적재) — 사람이 검토(승인/수정/반려)한 행과
    이번 달 이전(지난달까지) 행은 재스캔해도 덮어쓰지 않고 보존(잠금)한다.
    반환: (건수, pending수, TM-NO없어 제외된 건수, errors)."""
    wb = load_workbook(path, data_only=True)
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.worksheets[0]
    headers = [str(ws.cell(row=1, column=c).value or "").strip() for c in range(1, ws.max_column + 1)]
    idx = {h: j + 1 for j, h in enumerate(headers)}
    need = ["반입예정일자", "TM-NO", "Part"] + _OUTSOURCE_DETAIL_COLS
    miss = [h for h in need if h not in idx]
    if miss:
        return 0, 0, 0, [f"Sheet1 헤더 누락: {', '.join(miss)}"]

    route_cache = {}

    def has_jeonghyeong(tm):
        if tm not in route_cache:
            route_cache[tm] = conn.execute(
                "SELECT 1 FROM product_route WHERE tm_no=? AND process='정형'", (tm,)).fetchone() is not None
        return route_cache[tm]

    rows, pend, skip_no_tm = [], 0, 0
    for r in range(2, ws.max_row + 1):
        dv = ws.cell(row=r, column=idx["반입예정일자"]).value
        if dv in (None, ""):
            continue
        d = _d(dv)
        part = _OUTSOURCE_PART.get(str(ws.cell(row=r, column=idx["Part"]).value or "").strip().lower(), "")
        if not part:
            continue
        tmv = ws.cell(row=r, column=idx["TM-NO"]).value
        if isinstance(tmv, (datetime.date, datetime.datetime)):
            tm = ""                                   # TM-NO 날짜변환 → 제품 미지정
        else:
            tm = base_tmno(tmv)
        if not tm:
            skip_no_tm += 1
            continue                                  # TM-NO 없는 행은 라우팅 판단 불가 → 집계 제외
        for hdr in _OUTSOURCE_DETAIL_COLS:
            qty = _int(ws.cell(row=r, column=idx[hdr]).value)
            if qty <= 0:
                continue
            if hdr in _OUTSOURCE_DETAIL_DIRECT:
                proc = _OUTSOURCE_DETAIL_DIRECT[hdr]
            else:
                proc = "정형" if has_jeonghyeong(tm) else "소결"
            status = "pending" if (quarantine_100 and qty >= 100) else "confirmed"
            if status == "pending":
                pend += 1
            rows.append((d, tm, hdr, qty, part, proc, "공정", status, ""))
    n, _skipped, _locked = _reingest_preserve_reviewed(conn, "outsource", rows)
    return n, pend, skip_no_tm, []


_PROD_RE = re.compile(r"^(1part|2part)_(\d{8})_(\d{8})$", re.IGNORECASE)
PROD_FRESHNESS_KEY = "prod_last_file_end"   # 생산량 "데이터 최종입력일"(파일명 종료일 기준) 캐시 키


def _weekdays_in_range(d1, d2):
    """[d1,d2] 사이 평일(월~금)만 날짜 리스트(ISO 문자열)로. 평일이 하나도 없으면(범위가
    전부 토·일) 데이터 유실을 막기 위해 전체 달력일로 폴백한다."""
    days, cur = [], d1
    while cur <= d2:
        if cur.weekday() < 5:                      # 0=월 ... 4=금
            days.append(cur.isoformat())
        cur += datetime.timedelta(days=1)
    if days:
        return days
    cur = d1
    while cur <= d2:
        days.append(cur.isoformat())
        cur += datetime.timedelta(days=1)
    return days


def ingest_production_xlsx(conn, path):
    """생산수량 xlsx(ERP 제품창고재고금액조회) 적재.
    파일명 {1part|2part}_{시작}_{종료}.xlsx → 파트, 대상일자는 시작~종료 범위 중 **평일(월~금)만
    균등 분배**(토·일 제외 — 예: 목~일 범위면 목·금 2일로 분해). 하루짜리 파일이면 그대로 100%.
    C열=규격=TM-NO, I열=입고수량=생산수량, J열=입고금액(원)→÷1000=천원.
    변형 TM-NO는 base로 합산. 반환: (건수, errors)."""
    stem = os.path.splitext(os.path.basename(path))[0]
    mo = _PROD_RE.match(stem)
    if not mo:
        return 0, [f"파일명 형식 오류: {os.path.basename(path)} (예: 1part_20260723_20260723.xlsx)"]
    part = "VMS PART" if mo.group(1).lower() == "1part" else "TM PART"
    ymd1, ymd2 = mo.group(2), mo.group(3)
    d1 = datetime.date(int(ymd1[:4]), int(ymd1[4:6]), int(ymd1[6:8]))
    d2 = datetime.date(int(ymd2[:4]), int(ymd2[4:6]), int(ymd2[6:8]))
    target_days = _weekdays_in_range(d1, d2)

    # 대시보드 "데이터 최종입력일" 참고용: 실제 반영일(주말 제외로 앞당겨질 수 있음)이 아니라
    # 파일명에 적힌 종료일 자체를 기준으로 삼는다(사용자 확정).
    prev = db.get_setting(conn, PROD_FRESHNESS_KEY, "")
    if not prev or d2.isoformat() > prev:
        db.set_setting(conn, PROD_FRESHNESS_KEY, d2.isoformat())

    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    # 헤더행 탐색(규격·입고수량·입고금액 포함)
    hr = None
    for r in range(1, min(ws.max_row, 8) + 1):
        vals = [str(ws.cell(row=r, column=c).value or "").strip() for c in range(1, ws.max_column + 1)]
        if "규격" in vals and "입고수량" in vals:
            hr = r; headers = vals; break
    if hr is None:
        return 0, [f"{os.path.basename(path)}: 헤더행(규격/입고수량)을 찾을 수 없음"]
    c_tm = headers.index("규격") + 1
    c_qty = headers.index("입고수량") + 1
    c_amt = headers.index("입고금액") + 1

    agg = {}  # base_tm -> [qty, amt천원]
    for r in range(hr + 1, ws.max_row + 1):
        first = str(ws.cell(row=r, column=1).value or "").strip()
        if first == "TOTAL":
            continue
        raw = ws.cell(row=r, column=c_tm).value
        if isinstance(raw, (datetime.date, datetime.datetime)):
            continue  # 규격이 날짜면 TM-NO 아님 → 스킵
        tm = base_tmno(raw)
        if not tm:
            continue
        qty = _int(ws.cell(row=r, column=c_qty).value)
        if qty == 0:
            continue                                    # 반품 등 마이너스 수량은 순증감에 포함(스킵 X)
        amt = float(ws.cell(row=r, column=c_amt).value or 0) / 1000.0  # 원 → 천원
        cur = agg.setdefault(tm, [0, 0.0])
        cur[0] += qty; cur[1] += amt
    n_days = len(target_days)
    for tm, (qty, amt) in agg.items():
        qty_split = largest_remainder(qty, [1] * n_days) if n_days > 1 else [qty]
        amt_split = [amt / n_days] * n_days if n_days > 1 else [amt]
        for day, q, a in zip(target_days, qty_split, amt_split):
            conn.execute(
                "INSERT INTO production(d,tm_no,qty,amount,part) VALUES(?,?,?,?,?) "
                "ON CONFLICT(d,tm_no) DO UPDATE SET qty=excluded.qty, amount=excluded.amount, part=excluded.part",
                (day, tm, q, a, part))
    conn.commit()
    return len(agg), []


# 과거 KPI 실적표: 시트의 KPI명 → 내부 키
_HIST_KPI = {
    "COPQ": "copq_pct",
    "CUSTOMER INCIDENT": "incident",
    "WARRANTY": "warranty",
    "공정불량": "proc_ppm",
    "셋팅불량": "set_ppm",
    "SCRAP QUANTITY": "scrap_qty_pct",
    "SCRAP COST": "scrap_cost_pct",
}
_HIST_PART = {"1PART": "VMS PART", "2PART": "TM PART", "합계": "통합"}


def _hist_kpi_key(name):
    n = " ".join(str(name or "").split()).upper()
    if n in _HIST_KPI:
        return _HIST_KPI[n]
    for k, v in _HIST_KPI.items():          # 'Customer Incident Cost' 등 부분일치 허용
        if n.startswith(k):
            return v
    return None


def ingest_history_xlsx(conn, path):
    """과거 KPI 실적표 적재 (행=KPI×부문, 열=FY별 12개월).

    1행에 FY 라벨(FY26/FY27...)이 병합되어 있고, 2행이 월(4월~3월), 3행부터 데이터.
    A열=KPI, B열=부문(1PART/2PART/합계), C열=단위.
    반환: (건수, note[dict])
    """
    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]

    # 열 → (ym) 매핑: 1행 FY 라벨을 오른쪽으로 전파하며 2행 월과 조합
    colmap, fy = {}, None
    for c in range(4, ws.max_column + 1):
        v1 = str(ws.cell(row=1, column=c).value or "").strip()
        if v1.upper().startswith("FY"):
            try:
                fy = int(v1[2:])
            except ValueError:
                pass
        mon = str(ws.cell(row=2, column=c).value or "").strip().replace("월", "")
        if fy is None or not mon.isdigit():
            continue
        m = int(mon)
        year = 2000 + fy - 1 if m >= 4 else 2000 + fy      # FY는 4월~익년 3월
        colmap[c] = f"{year:04d}-{m:02d}"

    rows, skipped = [], []
    for r in range(3, ws.max_row + 1):
        kpi = _hist_kpi_key(ws.cell(row=r, column=1).value)
        part = _HIST_PART.get(str(ws.cell(row=r, column=2).value or "").strip().upper())
        unit = str(ws.cell(row=r, column=3).value or "").strip()
        if not kpi or not part:
            name = ws.cell(row=r, column=1).value
            if name:
                skipped.append(str(name))
            continue
        for c, ym in colmap.items():
            v = ws.cell(row=r, column=c).value
            if v in (None, ""):
                continue
            try:
                val = float(v)
            except (TypeError, ValueError):
                continue
            rows.append((ym, part, kpi, val, unit))
    conn.executemany(
        "INSERT INTO kpi_actual(ym,part,kpi,value,unit) VALUES(?,?,?,?,?) "
        "ON CONFLICT(ym,part,kpi) DO UPDATE SET value=excluded.value, unit=excluded.unit", rows)
    conn.commit()
    months = sorted({ym for ym, *_ in rows})
    return len(rows), {"기간": f"{months[0]}~{months[-1]}" if months else "-",
                       "지표": len({k for _y, _p, k, _v, _u in rows}),
                       "인식못한행": sorted(set(skipped))}


def ingest_svp(conn, path):
    _, rows = _rows(path)
    n = 0
    for i, r in rows:
        conn.execute("INSERT INTO svp(ym,part,amount) VALUES(?,?,?) "
                     "ON CONFLICT(ym,part) DO UPDATE SET amount=excluded.amount",
                     (_ym(r.get("년월")), str(r.get("파트", "")).strip(), float(r.get("금액(천원)") or 0)))
        n += 1
    conn.commit()
    return n




# ── 과거(Low_data/Low_Inventory) 이력 일괄 적재 ──────────
# Low_data: 월단위 사내불량 실적(공정/셋팅 별도, 1part/2part 별도).
# 열=불량유형(여러 공정 유형이 한 시트에 섞여 있음), 행=TM-NO별 실적.
# 발생공정은 컬럼이 아니라 **품명(B열) 문자열**로 결정한다.
import calendar

_HIST_DEFECT_RE = re.compile(r"^(1part|2part)_(공정|셋팅)불량_(\d{2})년\s*(\d{1,2})월$", re.IGNORECASE)
_HIST_INVENTORY_RE = re.compile(r"^(1part|2part)_(\d{1,2})월_입고수량$", re.IGNORECASE)
_HIST_PART = {"1part": "VMS PART", "2part": "TM PART"}
_HIST_PROC_KEYWORDS = ["성형", "정형", "압입", "가공"]   # 우선순위 순. 매칭 없으면 '후처리'
_HIST_IGNORE_COLS = {"폐기"}                              # 불량유형이 아닌 별도 집계라 무시


def _month_last_day(year, month):
    return f"{year:04d}-{month:02d}-{calendar.monthrange(year, month)[1]:02d}"


def _hist_process_of(name):
    n = str(name or "")
    for kw in _HIST_PROC_KEYWORDS:
        if kw in n:
            return kw
    return "후처리"


def ingest_history_defect_xlsx(conn, path, year_base=2000):
    """Low_data 월별 사내불량 실적 1개 파일 적재.

    파일명 {1part|2part}_{공정|셋팅}불량_{YY}년 {M}월.xlsx.
    행의 품명(B열)에 성형/정형/압입/가공 중 포함된 키워드로 그 행을 작성한 공정을 정하고
    (없으면 '후처리'), **성형 공정에서 작성한 행은 통째로 제외**한다(공정/셋팅 모두).
    나머지 행의 모든 불량유형 컬럼 수량은 그 시트 작성 공정 그대로 저장만 하고,
    최종 귀속 공정은 calc.Masters.resolve()가 마스터의 발생공정/배분기준을 보고 정한다.
    '폐기' 컬럼은 불량유형이 아니라서 무시한다. 날짜는 그 달 **말일**로 통일.
    마스터(defect_type)에 없는 불량명은 공백만 다르면 기존 이름에 맞추고,
    그래도 없으면 새 이름으로 마스터에 추가한다(process='', alloc_rule='').
    재실행 안전(batch_key로 기존 분 삭제 후 재적재).
    반환: (건수, note[dict]).
    """
    stem = os.path.splitext(os.path.basename(path))[0]
    mo = _HIST_DEFECT_RE.match(stem)
    if not mo:
        return 0, {"error": f"파일명 형식 오류: {os.path.basename(path)}"}
    part = _HIST_PART[mo.group(1).lower()]
    kind = mo.group(2)
    year = year_base + int(mo.group(3))
    month = int(mo.group(4))
    d = _month_last_day(year, month)
    batch_key = f"hist|{part}|{kind}|{year:04d}-{month:02d}"

    # 마스터 이름 정규화(공백 유무만 다른 표기 흡수): 공백 제거 후 비교
    master = {(r["part"], r["kind"]): r["name"]
              for r in conn.execute("SELECT part,kind,name FROM defect_type WHERE part=? AND kind=?",
                                    (part, kind))}
    by_nospace = {}
    for r in conn.execute("SELECT name FROM defect_type WHERE part=? AND kind=?", (part, kind)):
        by_nospace["".join(r["name"].split())] = r["name"]

    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    headers = [_clean_name(ws.cell(row=2, column=c).value) for c in range(1, ws.max_column + 1)]
    defect_cols = [(c + 1, headers[c]) for c in range(5, len(headers))
                  if headers[c] and headers[c] not in _HIST_IGNORE_COLS]

    new_types, rows = set(), []
    for r in range(3, ws.max_row + 1):
        first = str(ws.cell(row=r, column=1).value or "").strip()
        if first.upper() == "TOTAL" or not first:
            continue
        tm = base_tmno(ws.cell(row=r, column=3).value)
        if not tm:
            continue
        proc = _hist_process_of(ws.cell(row=r, column=2).value)
        if proc == "성형":
            continue                                    # 성형 시트 작성분은 D/B 미반영
        for c, raw_name in defect_cols:
            qty = _int(ws.cell(row=r, column=c).value)
            if qty <= 0:
                continue
            key = "".join(raw_name.split())
            name = by_nospace.get(key)
            if not name:
                name = raw_name
                by_nospace[key] = name
                new_types.add(name)                    # 마스터에 없는 이름 → 추가 대상
            rows.append((d, tm, name, qty, part, proc, kind, "direct", "confirmed", "", batch_key))

    if new_types:
        conn.executemany(
            "INSERT OR IGNORE INTO defect_type(part,kind,process,alloc_rule,name) "
            "VALUES(?,?,?,?,?)", [(part, kind, "", "", n) for n in new_types])
    conn.execute("DELETE FROM defect_entry WHERE batch_key=?", (batch_key,))
    conn.executemany(
        "INSERT INTO defect_entry(d,tm_no,defect_name,qty,part,process,kind,source,status,reg_user,batch_key) "
        "VALUES(?,?,?,?,?,?,?,?,?,?,?)", rows)
    conn.commit()
    return len(rows), {"기간": d, "신규불량유형": sorted(new_types)}


def ingest_history_inventory_xlsx(conn, path, year=2026):
    """Low_Inventory 월별 입고수량 1개 파일 적재 → production 테이블.
    파일명 {1part|2part}_{M}월_입고수량.xlsx. 컬럼은 기존 생산실적 xlsx와 동일
    (C=규격=TM-NO, I=입고수량, J=입고금액(원)→천원). 날짜는 그 달 **말일**.
    반환: (품목수, errors)."""
    stem = os.path.splitext(os.path.basename(path))[0]
    mo = _HIST_INVENTORY_RE.match(stem)
    if not mo:
        return 0, [f"파일명 형식 오류: {os.path.basename(path)}"]
    part = _HIST_PART[mo.group(1).lower()]
    month = int(mo.group(2))
    d = _month_last_day(year, month)

    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    hr = None
    for r in range(1, min(ws.max_row, 8) + 1):
        vals = [str(ws.cell(row=r, column=c).value or "").strip() for c in range(1, ws.max_column + 1)]
        if "규격" in vals and "입고수량" in vals:
            hr = r; headers = vals; break
    if hr is None:
        return 0, [f"{os.path.basename(path)}: 헤더행(규격/입고수량)을 찾을 수 없음"]
    c_tm = headers.index("규격") + 1
    c_qty = headers.index("입고수량") + 1
    c_amt = headers.index("입고금액") + 1

    agg = {}
    for r in range(hr + 1, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value or "").strip() == "TOTAL":
            continue
        raw = ws.cell(row=r, column=c_tm).value
        if isinstance(raw, (datetime.date, datetime.datetime)):
            continue
        tm = base_tmno(raw)
        if not tm:
            continue
        qty = _int(ws.cell(row=r, column=c_qty).value)
        if qty == 0:
            continue                                    # 반품 등 마이너스 수량은 순증감에 포함(스킵 X)
        amt = float(ws.cell(row=r, column=c_amt).value or 0) / 1000.0
        cur = agg.setdefault(tm, [0, 0.0])
        cur[0] += qty; cur[1] += amt
    for tm, (qty, amt) in agg.items():
        conn.execute(
            "INSERT INTO production(d,tm_no,qty,amount,part) VALUES(?,?,?,?,?) "
            "ON CONFLICT(d,tm_no) DO UPDATE SET qty=excluded.qty, amount=excluded.amount, part=excluded.part",
            (d, tm, qty, amt, part))
    conn.commit()
    return len(agg), []


def ingest_incident(conn, path):
    _, rows = _rows(path)
    n = 0
    for i, r in rows:
        conn.execute("INSERT INTO incident(d,part,customer,content) VALUES(?,?,?,?)",
                     (_d(r.get("일자")), str(r.get("파트", "")).strip(),
                      str(r.get("고객", "")).strip(), str(r.get("내용", "")).strip()))
        n += 1
    conn.commit()
    return n
