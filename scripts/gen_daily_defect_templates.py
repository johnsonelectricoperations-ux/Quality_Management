# -*- coding: utf-8 -*-
"""사내 불량 일일 입력 양식 생성기.

파일 1개 = 파트 × 공정 × 구분(공정/셋팅) × 하루.
  - 행: TM-NO (입력자가 직접 기입)
  - 열: 불량유형 (해당 파트·구분의 전체 유형, 모든 공정 시트 공통)
  - 셀: 수량
파일명 형식:  {PART}_{공정}_{구분}불량_{YYYYMMDD}.xlsx   (예: 1PART_성형_공정불량_20260724.xlsx)

열(불량유형) 목록은 templates/불량유형마스터_양식.xlsx (DATA_1PART/DATA_2PART)에서
구분(공정/셋팅) 기준으로 읽어온다. 마스터가 바뀌면 이 스크립트만 다시 돌리면 반영된다.

사용:  python scripts/gen_daily_defect_templates.py [YYYYMMDD] [출력폴더]
       (인자 없으면 오늘 날짜, templates/사내불량_양식/ 에 생성)
"""
import os
import sys
import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

MASTER = os.path.join("templates", "불량유형마스터_양식.xlsx")

# 파트별 공정 목록 (좌→우 공정 순서)
PART_PROCS = {
    "1PART": ["성형", "소결", "정형", "가공", "압입", "밴딩", "선별공정"],
    "2PART": ["성형", "소결", "정형", "가공", "선별공정"],
}
# 공정별로 존재하는 구분(불량 종류). 셋팅불량은 성형·정형·가공에만 있음.
PROC_KINDS = {
    "성형": ["공정", "셋팅"],
    "소결": ["공정"],
    "정형": ["공정", "셋팅"],
    "가공": ["공정", "셋팅"],
    "압입": ["공정"],
    "밴딩": ["공정"],
    "선별공정": ["공정"],
}

INPUT_ROWS = 100  # 입력용 빈 행 수

# 스타일
TITLE_FONT = Font(bold=True, size=13, color="FFFFFF")
TITLE_FILL = PatternFill("solid", fgColor="F26100")      # Johnson Orange
META_FONT = Font(bold=True, size=10)
HDR_FONT = Font(bold=True, size=9, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="4A5568")        # graphite
KEY_FILL = PatternFill("solid", fgColor="FBE5D6")        # TM-NO/품명 강조(연주황)
SUM_FONT = Font(bold=True, size=10)
SUM_FILL = PatternFill("solid", fgColor="EDF2F7")
THIN = Side(style="thin", color="CBD5E0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _clean(name):
    """헤더용: 개행 제거·공백 정리."""
    return " ".join(str(name).replace("\n", " ").split())


def defect_types(master_path):
    """{'1PART': {'공정':[...], '셋팅':[...]}, '2PART': {...}}  (마스터 등장 순서 유지)."""
    wb = load_workbook(master_path, data_only=True)
    out = {}
    for part, sheet in (("1PART", "DATA_1PART"), ("2PART", "DATA_2PART")):
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        kinds = {"공정": [], "셋팅": []}
        for r in rows[1:]:
            if all(c is None or str(c).strip() == "" for c in r):
                continue
            gubun = (r[0] or "").strip()
            name = _clean(r[2] or "")
            if gubun in kinds and name:
                kinds[gubun].append(name)
        out[part] = kinds
    return out


def build_one(part, proc, kind, cols, date_str, out_dir):
    """양식 파일 1개 생성."""
    kind_label = "공정불량" if kind == "공정" else "셋팅불량"
    wb = Workbook()
    ws = wb.active
    ws.title = "DATA"

    n_fixed = 3                       # No, TM-NO, 품명
    total_cols = n_fixed + len(cols)
    last_col = get_column_letter(total_cols)

    # 1행: 제목
    ws.merge_cells(f"A1:{last_col}1")
    c = ws["A1"]
    c.value = "사내 불량 일일 입력"
    c.font = TITLE_FONT
    c.fill = TITLE_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    # 2행: 메타 (파트/공정/구분/일자)
    ws["A2"] = "파트"; ws["B2"] = part
    ws["C2"] = "공정"; ws["D2"] = proc
    ws["E2"] = "구분"; ws["F2"] = kind_label
    ws["G2"] = "일자"; ws["H2"] = date_str
    for cell in ("A2", "C2", "E2", "G2"):
        ws[cell].font = META_FONT
    ws.row_dimensions[2].height = 18

    # 4행: 헤더
    hdr_row = 4
    headers = ["No", "TM-NO", "품명"] + cols
    for j, h in enumerate(headers, start=1):
        cell = ws.cell(row=hdr_row, column=j, value=h)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="bottom",
                                   wrap_text=True, textRotation=90 if j > n_fixed else 0)
    ws.row_dimensions[hdr_row].height = 110

    # 데이터 빈 행
    first_data = hdr_row + 1
    for i in range(INPUT_ROWS):
        r = first_data + i
        ws.cell(row=r, column=1, value=i + 1)          # No
        for j in range(1, total_cols + 1):
            cell = ws.cell(row=r, column=j)
            cell.border = BORDER
            if j <= 2:
                cell.fill = KEY_FILL                    # No/TM-NO 강조
            elif j == 3:
                cell.fill = KEY_FILL

    # 합계 행 (열별 SUM)
    sum_row = first_data + INPUT_ROWS
    ws.cell(row=sum_row, column=2, value="합계").font = SUM_FONT
    ws.cell(row=sum_row, column=2).fill = SUM_FILL
    for j in range(n_fixed + 1, total_cols + 1):
        col = get_column_letter(j)
        cell = ws.cell(row=sum_row, column=j,
                       value=f"=SUM({col}{first_data}:{col}{sum_row-1})")
        cell.font = SUM_FONT
        cell.fill = SUM_FILL
        cell.border = BORDER
    for j in range(1, total_cols + 1):
        ws.cell(row=sum_row, column=j).border = BORDER

    # 열 너비
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 18
    for j in range(n_fixed + 1, total_cols + 1):
        ws.column_dimensions[get_column_letter(j)].width = 6

    # 틀 고정: 헤더 + 앞 3열
    ws.freeze_panes = ws.cell(row=first_data, column=n_fixed + 1)

    fname = f"{part}_{proc}_{kind_label}_{date_str}.xlsx"
    path = os.path.join(out_dir, fname)
    wb.save(path)
    return fname, len(cols)


def main():
    date_str = sys.argv[1] if len(sys.argv) > 1 else datetime.date.today().strftime("%Y%m%d")
    out_dir = sys.argv[2] if len(sys.argv) > 2 else os.path.join("templates", "사내불량_양식")
    os.makedirs(out_dir, exist_ok=True)

    cols_by = defect_types(MASTER)
    made = []
    for part, procs in PART_PROCS.items():
        for proc in procs:
            for kind in PROC_KINDS[proc]:
                cols = cols_by[part][kind]
                fname, n = build_one(part, proc, kind, cols, date_str, out_dir)
                made.append((fname, n))
    print(f"[생성 완료] {len(made)}개 → {out_dir}")
    for fname, n in made:
        print(f"  - {fname}  (불량유형 {n}열)")


if __name__ == "__main__":
    main()
